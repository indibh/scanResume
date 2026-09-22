"""Extract text from supported resume document formats."""

from io import BytesIO
from pathlib import Path

from pypdf import PdfReader
from docx import Document


class ExtractionError(Exception):
    """Raised when a supported document cannot be converted to text."""


def extract_text(filename: str, content: bytes) -> str:
    """Return normalized text for a PDF, legacy Word, or DOCX file."""
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".pdf":
            text = _extract_pdf(content)
        elif extension == ".docx":
            text = _extract_docx(content)
        elif extension == ".doc":
            text = _extract_doc(content)
        else:
            raise ExtractionError("Unsupported file type.")
    except ExtractionError:
        raise
    except Exception as error:
        raise ExtractionError(f"Could not read the document: {error}") from error

    normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not normalized:
        raise ExtractionError(
            "No text was found. Image-only PDFs require OCR and are not supported."
        )
    return normalized


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        paragraphs.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(paragraphs)


def _extract_doc(content: bytes) -> str:
    """Use textract for the binary legacy Word format."""
    try:
        import textract
    except ImportError as error:
        raise ExtractionError(
            "Legacy .doc support requires the textract dependency. "
            "Install dependencies from scanResume/requirements.txt."
        ) from error

    try:
        return textract.process(None, input=content, extension="doc").decode(
            "utf-8", errors="replace"
        )
    except Exception as error:
        raise ExtractionError(f"Could not extract text from the .doc file: {error}") from error
